"""PPTX generation for handoff presentations"""
from datetime import datetime
from typing import Dict, Any, List
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from app.utils.errors import ExternalServiceError


class PPTXGenerator:
    """Generate PowerPoint presentations for ICU handoff"""
    
    def __init__(self):
        """Initialize PPTX generator"""
        pass
    
    def generate_handoff_deck(self, note_json: Dict[str, Any], 
                              clinical_summary: Dict[str, Any],
                              patient_id: str) -> Presentation:
        """
        Generate a 2-3 slide handoff presentation.
        
        Args:
            note_json: Structured note JSON from LLM
            clinical_summary: Clinical summary data
            patient_id: Patient identifier
            
        Returns:
            Presentation object
        """
        prs = Presentation()
        prs.slide_width = Inches(10)
        prs.slide_height = Inches(7.5)
        
        # Slide 1: Snapshot
        self._add_snapshot_slide(prs, note_json, clinical_summary, patient_id)
        
        # Slide 2: Systems Summary
        self._add_systems_slide(prs, note_json, clinical_summary)
        
        # Slide 3: To-do & Watchouts (if content exists)
        if note_json.get('todo') or note_json.get('watchouts'):
            self._add_todo_slide(prs, note_json)
        
        return prs
    
    def _add_snapshot_slide(self, prs: Presentation, note_json: Dict[str, Any],
                           clinical_summary: Dict[str, Any], patient_id: str):
        """Add snapshot slide"""
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
        
        # Title
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
        title_frame = title_box.text_frame
        title_frame.text = f"ICU Handoff - Patient {patient_id}"
        title_para = title_frame.paragraphs[0]
        title_para.font.size = Pt(24)
        title_para.font.bold = True
        
        # One-liner
        oneliner = note_json.get('one_liner', 'Not available')
        oneliner_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(9), Inches(0.8))
        oneliner_frame = oneliner_box.text_frame
        oneliner_frame.text = f"Status: {oneliner}"
        oneliner_para = oneliner_frame.paragraphs[0]
        oneliner_para.font.size = Pt(14)
        
        # Key vitals (if available)
        y_pos = 2.2
        vitals_text = []
        
        if clinical_summary.get('heart_rate'):
            hr = clinical_summary['heart_rate']
            if hr.get('last'):
                vitals_text.append(f"HR: {hr['last']:.0f} {hr.get('unit', '')}")
        
        if clinical_summary.get('mean_arterial_pressure'):
            map_val = clinical_summary['mean_arterial_pressure']
            if map_val.get('last'):
                vitals_text.append(f"MAP: {map_val['last']:.0f} {map_val.get('unit', '')}")
        
        if clinical_summary.get('oxygen_saturation'):
            spo2 = clinical_summary['oxygen_saturation']
            if spo2.get('last'):
                vitals_text.append(f"SpO2: {spo2['last']:.0f}%")
        
        if clinical_summary.get('fio2'):
            fio2 = clinical_summary['fio2']
            if fio2.get('last'):
                vitals_text.append(f"FiO2: {fio2['last']:.0f}%")
        
        if vitals_text:
            vitals_box = slide.shapes.add_textbox(Inches(0.5), Inches(y_pos), Inches(9), Inches(0.6))
            vitals_frame = vitals_box.text_frame
            vitals_frame.text = "Key Vitals: " + " | ".join(vitals_text)
            vitals_para = vitals_frame.paragraphs[0]
            vitals_para.font.size = Pt(12)
            y_pos += 0.8
        
        # Code status
        code_status = note_json.get('code_status', 'Not specified')
        code_box = slide.shapes.add_textbox(Inches(0.5), Inches(y_pos), Inches(9), Inches(0.5))
        code_frame = code_box.text_frame
        code_frame.text = f"Code Status: {code_status}"
        code_para = code_frame.paragraphs[0]
        code_para.font.size = Pt(12)
    
    def _add_systems_slide(self, prs: Presentation, note_json: Dict[str, Any],
                          clinical_summary: Dict[str, Any]):
        """Add systems summary slide"""
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        
        # Title
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
        title_frame = title_box.text_frame
        title_frame.text = "Systems Summary"
        title_para = title_frame.paragraphs[0]
        title_para.font.size = Pt(24)
        title_para.font.bold = True
        
        # Problem list (main content)
        problems = note_json.get('problem_list', [])
        y_pos = 1.2
        
        for i, problem in enumerate(problems[:5]):  # Limit to 5 problems
            problem_text = problem.get('problem', 'Unknown')
            assessment = problem.get('assessment', '')
            plan = problem.get('plan', '')
            
            # Problem header
            prob_box = slide.shapes.add_textbox(Inches(0.5), Inches(y_pos), Inches(9), Inches(0.4))
            prob_frame = prob_box.text_frame
            prob_frame.text = f"{i+1}. {problem_text}"
            prob_para = prob_frame.paragraphs[0]
            prob_para.font.size = Pt(14)
            prob_para.font.bold = True
            y_pos += 0.5
            
            # Assessment and plan (concise)
            if assessment or plan:
                detail_text = []
                if assessment:
                    detail_text.append(f"Assessment: {assessment[:100]}")  # Truncate
                if plan:
                    detail_text.append(f"Plan: {plan[:100]}")
                
                detail_box = slide.shapes.add_textbox(Inches(0.8), Inches(y_pos), Inches(8.5), Inches(0.5))
                detail_frame = detail_box.text_frame
                detail_frame.text = " | ".join(detail_text)
                detail_para = detail_frame.paragraphs[0]
                detail_para.font.size = Pt(11)
                y_pos += 0.7
            
            if y_pos > 6.5:  # Prevent overflow
                break
    
    def _add_todo_slide(self, prs: Presentation, note_json: Dict[str, Any]):
        """Add to-do and watchouts slide"""
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        
        # Title
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
        title_frame = title_box.text_frame
        title_frame.text = "To-Do & Watchouts"
        title_para = title_frame.paragraphs[0]
        title_para.font.size = Pt(24)
        title_para.font.bold = True
        
        y_pos = 1.2
        
        # To-do items
        todos = note_json.get('todo', [])
        if todos:
            todo_box = slide.shapes.add_textbox(Inches(0.5), Inches(y_pos), Inches(4.5), Inches(0.5))
            todo_frame = todo_box.text_frame
            todo_frame.text = "To-Do:"
            todo_para = todo_frame.paragraphs[0]
            todo_para.font.size = Pt(16)
            todo_para.font.bold = True
            y_pos += 0.6
            
            for todo in todos[:8]:  # Limit items
                item_box = slide.shapes.add_textbox(Inches(0.8), Inches(y_pos), Inches(4.2), Inches(0.4))
                item_frame = item_box.text_frame
                item_frame.text = f"• {todo[:80]}"  # Truncate long items
                item_para = item_frame.paragraphs[0]
                item_para.font.size = Pt(11)
                y_pos += 0.4
                if y_pos > 6.5:
                    break
        
        # Watchouts
        watchouts = note_json.get('watchouts', [])
        if watchouts:
            y_pos = 1.2
            watch_box = slide.shapes.add_textbox(Inches(5.5), Inches(y_pos), Inches(4.5), Inches(0.5))
            watch_frame = watch_box.text_frame
            watch_frame.text = "Watchouts:"
            watch_para = watch_frame.paragraphs[0]
            watch_para.font.size = Pt(16)
            watch_para.font.bold = True
            y_pos += 0.6
            
            for watch in watchouts[:8]:
                item_box = slide.shapes.add_textbox(Inches(5.8), Inches(y_pos), Inches(4.2), Inches(0.4))
                item_frame = item_box.text_frame
                item_frame.text = f"• {watch[:80]}"
                item_para = item_frame.paragraphs[0]
                item_para.font.size = Pt(11)
                y_pos += 0.4
                if y_pos > 6.5:
                    break
