"""Generate PDF, DOCX, and PNG copies of the text samples.

Run from the service root:

    py scripts/generate_samples.py
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document

SAMPLES_DIR = Path(__file__).resolve().parents[1] / "samples"
SOURCE_TXT = SAMPLES_DIR / "jane_doe_lab_result.txt"


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def build_pdf_bytes(text: str) -> bytes:
    """Build a single-page text PDF with an extractable text layer (no reportlab)."""
    lines = [line if line.strip() else " " for line in text.replace("\r\n", "\n").split("\n")]
    content_lines = ["BT", "/F1 11 Tf", "36 760 Td"]
    for index, line in enumerate(lines):
        escaped = _pdf_escape(line)
        if index:
            content_lines.append("0 -14 Td")
        content_lines.append(f"({escaped}) Tj")
    content_lines.append("ET")
    stream_content = "\n".join(content_lines).encode("latin-1", errors="replace")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Courier >>",
        b"<< /Length " + str(len(stream_content)).encode() + b" >>\nstream\n"
        + stream_content
        + b"\nendstream",
    ]

    buffer = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for body in objects:
        offsets.append(len(buffer))
        buffer += f"{len(offsets) - 1} 0 obj\n".encode() + body + b"\nendobj\n"

    xref_offset = len(buffer)
    count = len(objects) + 1
    buffer += f"xref\n0 {count}\n".encode()
    buffer += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        buffer += f"{offset:010d} 00000 n \n".encode()
    buffer += (
        f"trailer\n<< /Size {count} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF"
    ).encode()
    return bytes(buffer)


def build_docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_heading("Laboratory Report", level=1)
    for line in text.replace("\r\n", "\n").split("\n"):
        document.add_paragraph(line)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _load_font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        Path(r"C:\Windows\Fonts\consola.ttf"),
        Path(r"C:\Windows\Fonts\cour.ttf"),
        Path(r"C:\Windows\Fonts\arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"),
        Path("/usr/share/fonts/truetype/liberation/LiberationMono-Regular.ttf"),
    ]
    for path in candidates:
        if path.is_file():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def build_png_bytes(text: str) -> bytes:
    lines = text.replace("\r\n", "\n").split("\n")
    font = _load_font(20)
    padding = 32
    line_height = 28
    width = 980
    height = padding * 2 + max(line_height * (len(lines) + 1), 200)
    image = Image.new("RGB", (width, height), color=(255, 255, 255))
    draw = ImageDraw.Draw(image)
    y = padding
    for line in lines:
        draw.text((padding, y), line, fill=(20, 20, 20), font=font)
        y += line_height
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def main() -> None:
    text = SOURCE_TXT.read_text(encoding="utf-8")
    outputs = {
        SAMPLES_DIR / "jane_doe_lab_result.pdf": build_pdf_bytes(text),
        SAMPLES_DIR / "jane_doe_lab_result.docx": build_docx_bytes(text),
        SAMPLES_DIR / "jane_doe_lab_result.png": build_png_bytes(text),
        SAMPLES_DIR / "jane_doe_old_lab_result.pdf": build_pdf_bytes(
            (SAMPLES_DIR / "jane_doe_old_lab_result.txt").read_text(encoding="utf-8")
        ),
        SAMPLES_DIR / "john_smith_lab_result.pdf": build_pdf_bytes(
            (SAMPLES_DIR / "john_smith_lab_result.txt").read_text(encoding="utf-8")
        ),
    }
    for path, data in outputs.items():
        path.write_bytes(data)
        print(f"Wrote {path.name} ({len(data)} bytes)")


if __name__ == "__main__":
    main()
