"""DOCX text extraction (new dependency: python-docx -- no existing lib in this repo
handles Word documents)."""
from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

MIN_EXTRACTED_TEXT_CHARS = 1


def extract_docx_text(file_bytes: bytes) -> str:
    """Extract paragraph and table text from a DOCX file.

    Raises:
        ValueError: If the file is corrupted/unreadable or not a valid DOCX package.
    """
    try:
        document = Document(BytesIO(file_bytes))
    except PackageNotFoundError as exc:
        raise ValueError("Unreadable or corrupted DOCX file") from exc
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"Unreadable or corrupted DOCX file: {exc}") from exc

    lines: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            lines.append(paragraph.text.strip())

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    text = "\n".join(lines)
    if len(text.strip()) < MIN_EXTRACTED_TEXT_CHARS:
        raise ValueError("No extractable text found in DOCX file")
    return text
