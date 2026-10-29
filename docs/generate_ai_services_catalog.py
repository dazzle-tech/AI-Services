"""Generate docs/AI-Services-Catalog.docx from the curated service list."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "AI-Services-Catalog.docx"

# folder, display name, port, category, purpose
SERVICES = [
    (
        "PatientTimelineService",
        "Patient Timeline Service",
        "8001",
        "Clinical Documentation & Summaries",
        "Builds a chronological timeline of clinically important events from structured and semi-structured patient data.",
    ),
    (
        "NurseTaskPrioritizationService",
        "Nurse Task Prioritization Service",
        "8002",
        "Nursing & Care Coordination",
        "Prioritizes nursing tasks from monitoring data, med schedules, abnormal findings, and operational context.",
    ),
    (
        "LabResultInterpreterService",
        "Lab Result Interpreter Service",
        "8003",
        "Clinical Decision Support & Safety",
        "Interprets lab values with optional patient context: abnormal findings, trends, patterns, and follow-up suggestions.",
    ),
    (
        "SmartDischargePlannerService",
        "Smart Discharge Planner Service",
        "8004",
        "Discharge & Continuity",
        "Assess discharge readiness, blockers, follow-up needs, med reconciliation concerns, and a draft discharge plan.",
    ),
    (
        "DischargeReportGenerator",
        "Discharge Report Generator",
        "8005",
        "Discharge & Continuity",
        "Generates professional discharge summaries from clinical documentation using GPT-4.",
    ),
    (
        "MedicationTestOrdersValidation",
        "Medication & Test Orders Validation",
        "8006",
        "Clinical Decision Support & Safety",
        "Validates medications and diagnostic test orders against patient data for safety, contraindications, and conflicts.",
    ),
    (
        "RecommendationsService",
        "Clinical Recommendations Service",
        "8007",
        "Clinical Decision Support & Safety",
        "Produces evidence-based, actionable clinical recommendations from full patient context.",
    ),
    (
        "MedicalGuidelineValidation",
        "Medical Guideline Validation",
        "8008",
        "Clinical Decision Support & Safety",
        "Validates active medical orders against evidence-based clinical guidelines.",
    ),
    (
        "SummarizationService",
        "Clinical Summary Service",
        "8009",
        "Clinical Documentation & Summaries",
        "Turns structured patient data into concise professional clinical summaries while preserving facts.",
    ),
    (
        "AIAutoPopulation",
        "AI Auto-Population Service",
        "8010",
        "Clinical Documentation & Summaries",
        "Extracts structured medical data from clinician free text and cross-checks it against the patient record.",
    ),
    (
        "QualityDischargeReport",
        "Discharge QA Service",
        "8011",
        "Discharge & Continuity",
        "QA engine that validates discharge reports against records, onsite documentation, templates, and quality rules.",
    ),
    (
        "OCRParsingService",
        "OCR Parsing Service",
        "8012",
        "Documents, OCR & Coding",
        "Extracts text from images via OCR and parses it into structured fields with OpenAI.",
    ),
    (
        "IcuSummarizer",
        "ICU Summarizer",
        "8013",
        "Clinical Documentation & Summaries",
        "Generates ICU daily summaries and handoff presentations from structured flowsheet data.",
    ),
    (
        "SpecialistAlertService",
        "Specialist Alert Service",
        "8014",
        "Clinical Decision Support & Safety",
        "Returns grounded alerts for specialist consult triggers, urgent review, and medication/safety concerns.",
    ),
    (
        "MedicalImageInterpretationAssistService",
        "Medical Image Interpretation Assist (X-ray MVP)",
        "8015",
        "Radiology & Imaging",
        "Chest X-ray assist only: returns suspected/possible AI candidate findings; never confirms a diagnosis.",
    ),
    (
        "RadiologyImageQaAI",
        "Radiology Image QC AI (MVP)",
        "8016",
        "Radiology & Imaging",
        "Technical image quality control for radiology; initial focus CT abdomen/pelvis anatomical coverage QC.",
    ),
    (
        "RadiolgyWorkFlow",
        "Radiology AI Pipeline",
        "8090",
        "Radiology & Imaging",
        "Orchestrates radiology AI steps (QC 8016, interpretation 8015, report filling 8024, autofill 8022).",
    ),
    (
        "Chatbot",
        "Clinician Chatbot (MedAI Orchestrator)",
        "8017",
        "Chatbots & Assistants",
        "Clinician-facing multi-service chatbot orchestrator (SQL gen 8018, validator 8019, formatter 8020, UI 8080).",
    ),
    (
        "MedicalImageTemplateAutofill",
        "Radiology Template Selection & Autofill",
        "8022",
        "Radiology & Imaging",
        "Two-stage service: RAG template selection then autofill for CT/MRI/US/X-ray/mammography report drafting.",
    ),
    (
        "SepsisEarlyDetection",
        "Sepsis Early Detection (SepsisSentinel)",
        "8023",
        "Clinical Decision Support & Safety",
        "Analyses 24-hour patient time-series data and returns AI-powered early sepsis risk signals.",
    ),
    (
        "RadiologyReportFilling",
        "Medical Imaging Assist / Radiology Report Filling",
        "8024",
        "Radiology & Imaging",
        "Validates DICOM metadata, corrects reporting errors, and reconciles AI image analysis with clinical reporting.",
    ),
    (
        "BillingCoder",
        "CodingAssist (Billing Coder)",
        "8025",
        "Documents, OCR & Coding",
        "Auto-coding and charge capture: ICD-10-CM / CPT / HCPCS from clinical documentation with compliance checks.",
    ),
    (
        "ConvoVoiceAgent",
        "ConvoVoiceAgent",
        "8036",
        "Voice & Transcription",
        "Single integration for doctor–patient recordings; orchestrates ConvoScribe (8026) and SystemDisplayPlugin (8031).",
    ),
    (
        "STT_Service",
        "STT Service (Radiology Speech-to-Text)",
        "8027",
        "Voice & Transcription",
        "Local Whisper-based STT microservice specialized for radiology dictation.",
    ),
    (
        "NurseHandOver",
        "Nurse HandOver (SBAR)",
        "8028",
        "Nursing & Care Coordination",
        "Stateless SBAR handoff per patient: Situation / Background / Assessment / Recommendation with priority.",
    ),
    (
        "MedicalDocumentProcessor",
        "Medical Document Processor",
        "8029",
        "Documents, OCR & Coding",
        "Intake pipeline for uploaded medical documents: identity check, relevance, translation, typed extraction.",
    ),
    (
        "ORVoiceAgent",
        "OR Voice Agent",
        "8035",
        "Voice & Transcription",
        "EMR Record-button integration for OR nursing tabs; routes voice to ORScribe (8030) and ORDisplayPlugin (8032).",
    ),
    (
        "LiveVisitScribe",
        "LiveVisitScribe",
        "8037",
        "Voice & Transcription",
        "Live visit WebSocket transcript (gpt-live-transcribe) with doctor/patient lines and a picked-up chart scratch list.",
    ),
    (
        "ClinicalSummaryVoice",
        "Clinical Summary Voice",
        "8039",
        "Voice & Transcription",
        "One-speaker clinical summary dictation: transcribe audio and return a summary (not multi-speaker SOAP).",
    ),
    (
        "chatbot-patient",
        "Patient Chatbot (MedAI Patient Assistant)",
        "8038",
        "Chatbots & Assistants",
        "Patient-facing agent with separate prompts and tools from the clinician chatbot; no shared runtime.",
    ),
]

CATEGORIES = [
    "Voice & Transcription",
    "Clinical Documentation & Summaries",
    "Clinical Decision Support & Safety",
    "Nursing & Care Coordination",
    "Radiology & Imaging",
    "Documents, OCR & Coding",
    "Chatbots & Assistants",
    "Discharge & Continuity",
]

RELATED = [
    (
        "ConvoScribe",
        "ConvoVoiceAgent/services/ConvoScribe",
        "Port 8026 — transcription + clinical note worker used by ConvoVoiceAgent.",
    ),
    (
        "SystemDisplayPlugin",
        "ConvoVoiceAgent/services/SystemDisplayPlugin",
        "Port 8031 — display/plugin surface for ConvoVoiceAgent results.",
    ),
    (
        "ORScribe",
        "ORVoiceAgent/services/ORScribe",
        "Port 8030 — OR voice transcription/extraction worker for ORVoiceAgent.",
    ),
    (
        "ORDisplayPlugin",
        "ORVoiceAgent/services/ORDisplayPlugin",
        "Port 8032 — display/plugin paths mirrored by ORVoiceAgent.",
    ),
]


def set_run_font(run, name="Calibri", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def add_heading_custom(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    return p


def main() -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        title.add_run("AI Services Catalog"),
        size=28,
        bold=True,
        color=RGBColor(0x1F, 0x3A, 0x5F),
    )

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        sub.add_run("Dazzle Tech · AI-Services repository"),
        size=12,
        color=RGBColor(0x55, 0x55, 0x55),
    )

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        meta.add_run(f"Generated {date.today().isoformat()}  ·  {len(SERVICES)} primary services"),
        size=10,
        color=RGBColor(0x66, 0x66, 0x66),
    )

    intro = doc.add_paragraph()
    set_run_font(
        intro.add_run(
            "This catalog lists the available AI microservices in the AI-Services workspace. "
            "Each entry includes the folder name, typical local port, category, and a short purpose. "
            "Ports come from each service’s defaults / .env.example and may differ in deployment."
        ),
        size=11,
    )

    add_heading_custom(doc, "1. Summary table", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["#", "Service", "Port", "Category"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                set_run_font(run, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade_cell(cell, "1F3A5F")

    for i, (_folder, name, port, cat, _purpose) in enumerate(SERVICES, 1):
        row = table.add_row().cells
        for j, v in enumerate([str(i), name, port, cat]):
            row[j].text = v
            for p in row[j].paragraphs:
                for run in p.runs:
                    set_run_font(run, size=9)
        if i % 2 == 0:
            for j in range(4):
                shade_cell(row[j], "F2F5F9")

    add_heading_custom(doc, "2. Services by category", level=1)
    by_cat = {c: [] for c in CATEGORIES}
    for s in SERVICES:
        by_cat[s[3]].append(s)

    for cat in CATEGORIES:
        items = by_cat[cat]
        if not items:
            continue
        add_heading_custom(doc, f"{cat} ({len(items)})", level=2)
        for folder, name, port, _cat, purpose in items:
            p = doc.add_paragraph()
            set_run_font(p.add_run(name), size=12, bold=True, color=RGBColor(0x1F, 0x3A, 0x5F))
            meta_p = doc.add_paragraph()
            set_run_font(
                meta_p.add_run(f"Folder: {folder}    ·    Port: {port}"),
                size=9,
                color=RGBColor(0x66, 0x66, 0x66),
            )
            body = doc.add_paragraph()
            set_run_font(body.add_run(purpose), size=11)

    add_heading_custom(doc, "3. Related nested components", level=1)
    p = doc.add_paragraph()
    set_run_font(
        p.add_run(
            "These are not standalone top-level deployables in the same way; they ship inside a parent agent."
        ),
        size=11,
    )
    rt = doc.add_table(rows=1, cols=3)
    rt.style = "Table Grid"
    for i, h in enumerate(["Component", "Location", "Role"]):
        cell = rt.rows[0].cells[i]
        cell.text = h
        for para in cell.paragraphs:
            for run in para.runs:
                set_run_font(run, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade_cell(cell, "1F3A5F")
    for name, loc, role in RELATED:
        row = rt.add_row().cells
        for j, v in enumerate([name, loc, role]):
            row[j].text = v
            for para in row[j].paragraphs:
                for run in para.runs:
                    set_run_font(run, size=9)

    add_heading_custom(doc, "4. Notes", level=1)
    notes = [
        "Auth for browser/EMR clients is expected at the API gateway; individual services focus on clinical AI work.",
        "Many services require an OpenAI-compatible API key via .env (never commit real .env files).",
        "Each top-level service and nested worker has a unique default listen port (LiveVisitScribe 8037, ClinicalSummaryVoice 8039).",
        "Chatbot is a multi-process stack (orchestrator + helpers + optional frontend), not a single FastAPI file.",
        "Folder spelling RadiolgyWorkFlow is intentional as checked into the repository.",
    ]
    for n in notes:
        bp = doc.add_paragraph(n, style="List Bullet")
        for run in bp.runs:
            set_run_font(run, size=11)

    doc.save(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size} bytes)")
    print(f"Services: {len(SERVICES)}")


if __name__ == "__main__":
    main()
